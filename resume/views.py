from django.shortcuts import render
import json
import os
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from docx import Document

from accounts.models import User
from .ats_scoring import build_structured_resume, compute_ats_score, extract_text

try:
    from google import genai
except ImportError:
    genai = None
from .models import JDMatchResult, JobDescription, Resume, ResumeAnalysis, ResumeVersion

allowed_extensions = {
    'pdf': 'pdf', 'docx': 'docx', 'doc': 'docx',
    # pictures are accepted only so we can tell the user they score very low
    'png': 'image', 'jpg': 'image', 'jpeg': 'image',
}
max_file_size = 5 * 1024 * 1024
min_text_length = 30


""" Anonymous visitors: resume is scored in-memory only, nothing is written
    to the database — matches the "No account needed to see your score" flow.
    Logged-in users (request.session['user_id']): the upload, its parsed
    text, and the analysis are persisted into resume / resume_version /
    resume_analysis (and job_descriptions / jd_match_results if a JD was
    pasted), per the data dictionary.
    Image-based resumes (scans, screenshots, jpg/png) get a very low score
    plus an explanation, and are not persisted. """


@require_http_methods(['POST'])
def download_optimized_resume(request):
    """Generate a professionally formatted DOCX from the edited resume JSON."""
    raw = request.POST.get('resume_json') or ''
    try:
        data = json.loads(raw)
    except (TypeError, ValueError):
        return JsonResponse({'error': 'Invalid resume editor data.'}, status=400)

    if not isinstance(data, dict):
        return JsonResponse({'error': 'Invalid resume structure.'}, status=400)

    if not optimized_data_to_plain_text(data).strip():
        return JsonResponse({'error': 'Resume content is empty.'}, status=400)

    try:
        document = build_professional_resume_docx(data)
    except Exception as exc:
        import traceback
        traceback.print_exc()
        return JsonResponse({'error': 'Could not create the DOCX resume.', 'details': str(exc)}, status=500)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="ATS_Optimized_Resume.docx"'
    document.save(response)
    return response


def _clean_ai_json(value):
    """Remove accidental markdown fences around Gemini JSON output."""
    if not isinstance(value, str):
        return value
    value = value.strip()
    if value.startswith('```'):
        value = value.split('\n', 1)[1] if '\n' in value else value
        if value.endswith('```'):
            value = value[:-3]
    return value.strip()


def _as_list(value):
    if isinstance(value, list):
        return value
    if value in (None, ''):
        return []
    return [value]


def _clean_string(value):
    return str(value or '').strip()


def _normalize_ai_resume(data):
    """Guarantee one stable structure for both AI output and the editor."""
    if not isinstance(data, dict):
        return None

    result = {
        'name': _clean_string(data.get('name')),
        'headline': _clean_string(data.get('headline')),
        'contact': [_clean_string(x) for x in _as_list(data.get('contact')) if _clean_string(x)],
        'summary': _clean_string(data.get('summary')),
        'skills': {},
        'experience': [],
        'projects': [],
        'education': [],
        'certifications': [],
        'achievements': [],
        'languages': [],
    }

    skills = data.get('skills') or {}
    if isinstance(skills, dict):
        for category, values in skills.items():
            category = _clean_string(category)
            if not category:
                continue
            cleaned = [_clean_string(x) for x in _as_list(values) if _clean_string(x)]
            if cleaned:
                result['skills'][category] = cleaned

    def normalize_entry(item, kind):
        if isinstance(item, str):
            return {'title': '', 'company': '', 'location': '', 'dates': '', 'technologies': [], 'details': [], 'bullets': [item.strip()]} if kind == 'experience' else (
                {'name': item.strip(), 'technologies': [], 'description': '', 'bullets': []} if kind == 'projects' else
                {'degree': item.strip(), 'institution': '', 'location': '', 'dates': '', 'details': []}
            )
        if not isinstance(item, dict):
            return None

        if kind == 'experience':
            return {
                'title': _clean_string(item.get('title') or item.get('role') or item.get('position')),
                'company': _clean_string(item.get('company') or item.get('organization')),
                'location': _clean_string(item.get('location')),
                'dates': _clean_string(item.get('dates') or item.get('date') or item.get('duration')),
                'bullets': [_clean_string(x) for x in _as_list(item.get('bullets') or item.get('responsibilities') or item.get('details')) if _clean_string(x)],
            }
        if kind == 'projects':
            return {
                'name': _clean_string(item.get('name') or item.get('title') or item.get('project')),
                'technologies': [_clean_string(x) for x in _as_list(item.get('technologies') or item.get('technology') or item.get('tech_stack')) if _clean_string(x)],
                'description': _clean_string(item.get('description')),
                'bullets': [_clean_string(x) for x in _as_list(item.get('bullets') or item.get('details')) if _clean_string(x)],
            }
        return {
            'degree': _clean_string(item.get('degree') or item.get('qualification') or item.get('title')),
            'institution': _clean_string(item.get('institution') or item.get('college') or item.get('university')),
            'location': _clean_string(item.get('location')),
            'dates': _clean_string(item.get('dates') or item.get('date') or item.get('duration')),
            'details': [_clean_string(x) for x in _as_list(item.get('details') or item.get('bullets')) if _clean_string(x)],
        }

    for kind in ('experience', 'projects', 'education'):
        for item in _as_list(data.get(kind)):
            normalized = normalize_entry(item, kind)
            if normalized and any(normalized.values()):
                result[kind].append(normalized)

    for key in ('certifications', 'achievements', 'languages'):
        for item in _as_list(data.get(key)):
            if isinstance(item, dict):
                value = item.get('name') or item.get('title') or item.get('text') or item.get('value')
            else:
                value = item
            value = _clean_string(value)
            if value:
                result[key].append(value)

    return result


def _fallback_structured_resume(resume_text):
    """Convert the existing parser output into the new structured editor shape."""
    old = build_structured_resume(resume_text)
    data = {
        'name': old.get('name', ''),
        'headline': old.get('headline', ''),
        'contact': old.get('contact', []),
        'summary': old.get('summary', ''),
        'skills': old.get('skills', {}),
        'experience': [], 'projects': [], 'education': [],
        'certifications': [], 'achievements': [], 'languages': [],
    }
    for item in old.get('experience', []):
        data['experience'].append({'title': '', 'company': '', 'location': '', 'dates': '', 'bullets': [_clean_string(item)]})
    for item in old.get('projects', []):
        data['projects'].append({'name': _clean_string(item), 'technologies': [], 'description': '', 'bullets': []})
    for item in old.get('education', []):
        data['education'].append({'degree': _clean_string(item), 'institution': '', 'location': '', 'dates': '', 'details': []})
    for key in ('certifications', 'achievements', 'languages'):
        data[key] = [_clean_string(x) for x in old.get(key, []) if _clean_string(x)]
    return data


def improve_resume_with_ai(resume_text, jd_text=""):

    fallback = build_structured_resume(resume_text)

    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        return fallback

    try:

        client = genai.Client(
            api_key=api_key
        )

        prompt = f"""
You are ResumeIQ's professional resume improvement engine.

Convert the source resume into a clean, professional,
ATS-friendly editable resume.

IMPORTANT:

1. Restore missing spaces between words.

Examples:

VISHAKHAPARESHKUMARPATEL
-> Vishakha Pareshkumar Patel

AspiringSoftwareDeveloper
-> Aspiring Software Developer

Final-yearMCAstudent
-> Final-year MCA student

ClinicManagementSystem
-> Clinic Management System

2. NEVER mix contact information with the summary.

Contact information must contain only:
- location
- phone
- email
- LinkedIn
- GitHub
- portfolio
- other professional profile links

3. PROFESSIONAL SUMMARY must contain only the candidate's
professional summary.

4. Preserve actual information from the source.

5. NEVER invent:
- company
- job
- degree
- project
- technology
- certification
- date
- achievement
- metric
- URL
- responsibility

6. You may improve:
- grammar
- sentence structure
- capitalization
- readability
- action verbs
- bullet-point wording

7. Do not add job-description skills unless the source resume
actually supports them.

8. Keep all real projects, education, certifications,
skills and achievements.

Return ONLY valid JSON.

JSON format:

{{
    "name": "",
    "headline": "",
    "contact": [],

    "summary": "",

    "skills": {{}},

    "experience": [
        {{
            "title": "",
            "company": "",
            "location": "",
            "dates": "",
            "bullets": []
        }}
    ],

    "projects": [
        {{
            "name": "",
            "technologies": [],
            "description": "",
            "bullets": []
        }}
    ],

    "education": [
        {{
            "degree": "",
            "institution": "",
            "location": "",
            "dates": "",
            "details": []
        }}
    ],

    "certifications": [],
    "achievements": [],
    "languages": []
}}

SOURCE RESUME:

{resume_text}

JOB DESCRIPTION:

{jd_text}
"""

        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt,
            config={
                "temperature": 0.2,
                "response_mime_type": "application/json",
            }
        )

        result = json.loads(
            response.text
        )

        return normalize_ai_resume(result)

    except Exception as exc:

        print(
            "Resume AI improvement failed:",
            exc
        )

        return fallback

def normalize_ai_resume(data):

    if not isinstance(data, dict):
        return {}

    result = {
        "name": str(
            data.get("name", "")
        ).strip(),

        "headline": str(
            data.get("headline", "")
        ).strip(),

        "contact": [],

        "summary": str(
            data.get("summary", "")
        ).strip(),

        "skills": {},

        "experience": [],

        "projects": [],

        "education": [],

        "certifications": [],

        "achievements": [],

        "languages": [],
    }

    # ---------------------------------------------------------
    # CONTACT
    # ---------------------------------------------------------

    contact = data.get(
        "contact",
        []
    )

    if isinstance(contact, str):
        contact = [
            x.strip()
            for x in contact.split("|")
            if x.strip()
        ]

    if isinstance(contact, list):

        result["contact"] = [
            str(x).strip()
            for x in contact
            if str(x).strip()
        ]

    # ---------------------------------------------------------
    # SKILLS
    # ---------------------------------------------------------

    skills = data.get(
        "skills",
        {}
    )

    if isinstance(skills, dict):

        for category, values in skills.items():

            if isinstance(values, str):

                values = [
                    x.strip()
                    for x in values.split(",")
                    if x.strip()
                ]

            if isinstance(values, list):

                result["skills"][
                    str(category).strip()
                ] = [
                    str(x).strip()
                    for x in values
                    if str(x).strip()
                ]

    # ---------------------------------------------------------
    # EXPERIENCE
    # ---------------------------------------------------------

    for item in data.get(
        "experience",
        []
    ):

        if not isinstance(item, dict):
            continue

        result["experience"].append({

            "title": str(
                item.get("title", "")
            ).strip(),

            "company": str(
                item.get("company", "")
            ).strip(),

            "location": str(
                item.get("location", "")
            ).strip(),

            "dates": str(
                item.get("dates", "")
            ).strip(),

            "bullets": [
                str(x).strip()
                for x in item.get(
                    "bullets",
                    []
                )
                if str(x).strip()
            ],
        })

    # ---------------------------------------------------------
    # PROJECTS
    # ---------------------------------------------------------

    for item in data.get(
        "projects",
        []
    ):

        if not isinstance(item, dict):
            continue

        technologies = item.get(
            "technologies",
            []
        )

        if isinstance(
            technologies,
            str
        ):

            technologies = [
                x.strip()
                for x in technologies.split(",")
                if x.strip()
            ]

        result["projects"].append({

            "name": str(
                item.get("name", "")
            ).strip(),

            "technologies": [
                str(x).strip()
                for x in technologies
                if str(x).strip()
            ],

            "description": str(
                item.get(
                    "description",
                    ""
                )
            ).strip(),

            "bullets": [
                str(x).strip()
                for x in item.get(
                    "bullets",
                    []
                )
                if str(x).strip()
            ],
        })

    # ---------------------------------------------------------
    # EDUCATION
    # ---------------------------------------------------------

    for item in data.get(
        "education",
        []
    ):

        if not isinstance(item, dict):
            continue

        result["education"].append({

            "degree": str(
                item.get("degree", "")
            ).strip(),

            "institution": str(
                item.get("institution", "")
            ).strip(),

            "location": str(
                item.get("location", "")
            ).strip(),

            "dates": str(
                item.get("dates", "")
            ).strip(),

            "details": [
                str(x).strip()
                for x in item.get(
                    "details",
                    []
                )
                if str(x).strip()
            ],
        })

    # ---------------------------------------------------------
    # SIMPLE LIST SECTIONS
    # ---------------------------------------------------------

    for key in [
        "certifications",
        "achievements",
        "languages",
    ]:

        values = data.get(
            key,
            []
        )

        if isinstance(
            values,
            str
        ):
            values = [
                values
            ]

        result[key] = [
            str(x).strip()
            for x in values
            if str(x).strip()
        ]

    return result

@require_http_methods(['POST'])
def analyze_resume(request):

    upload = request.FILES.get('resume')
    jd_text = (request.POST.get('jd') or '').strip()

    if not upload:
        return JsonResponse({"error": "Please attach a resume file."}, status=400)

    ext = upload.name.rsplit('.', 1)[-1].lower() if '.' in upload.name else ""

    file_type = allowed_extensions.get(ext)

    if not file_type:
        return JsonResponse({"error": 'Only PDF or DOCX files are supported.'}, status=400)

    if upload.size > max_file_size:
        return JsonResponse({"error": 'File is larger than 5MB.'}, status=400)

    raw_bytes = upload.read()

    """ try:
        resume_text, has_tables, content_meta = extract_text(raw_bytes, file_type)
    except Exception:
        return JsonResponse({"error": "We couldn't read that file. Try re-saving it and uploading again."}, status=422) """

    try:
        resume_text, has_tables, content_meta = extract_text(
            raw_bytes,
            file_type
        )

    except Exception as e:
        import traceback
        traceback.print_exc()

        return JsonResponse(
            {
                "error": "Resume extraction failed.",
                "details": str(e)
            },
            status=422
        )

    resume_text = resume_text or ''

    # Empty text is no longer a hard error when the file contains images:
    # that is exactly the "image resume" case and it should be scored (very low).
    has_images = content_meta.get('image_count', 0) > 0 or content_meta.get('is_image_file')
    if len(resume_text.strip()) < min_text_length and not has_images:
        return JsonResponse({"error": "We couldn't find readable text in that file."}, status=422)

    result = compute_ats_score(resume_text, jd_text or None, has_tables=has_tables, content_meta=content_meta)
    saved = False

    user_id = request.session.get('user_id')
    if user_id and not result['image_based']:
        try:
            user = User.objects.get(user_id=user_id)
        except User.DoesNotExist:
            user = None

        if user:
            stored_path = default_storage.save(f'resumes/{user.user_id}/{upload.name}', ContentFile(raw_bytes))

            resume = Resume.objects.create(
                user=user,
                file_type=file_type,
                file_path=stored_path,
                parsed_text=resume_text
            )

            next_version_number = (ResumeVersion.objects.filter(versions=resume).count()) + 1

            version = ResumeVersion.objects.create(
                versions=resume,
                version_number=next_version_number,
                content_snapshot=resume_text
            )

            ResumeAnalysis.objects.create(
                analyses=version,
                ats_score=result['ats_score'],
                grammar_issues=result['grammar_issues'],
                passive_voice_flags=result['passive_voice_flags'],
                missing_section=result['missing_sections']
            )

            if jd_text:
                jd = JobDescription.objects.create(job_descriptions=user, jd_text=jd_text)
                JDMatchResult.objects.create(
                    user=user,
                    version=version,
                    jd=jd,
                    match_percentage=result['keyword_match'],
                    missing_keywords=result['missing_keywords'],
                    search_type='with_resume'
                )

            saved = True

    return JsonResponse({
        'ats_score': result['ats_score'],
        'formatting_score': result['formatting_score'],
        'section_score': result['section_score'],
        'keyword_match': result['keyword_match'],
        'missing_sections': result['missing_sections'],
        'missing_keywords': result['missing_keywords'],
        'jd_gaps': result.get('jd_gaps', []),
        'optimized_resume': improve_resume_with_ai(resume_text, jd_text),
        'image_based': result['image_based'],
        'image_warnings': result['image_warnings'],
        'saved': saved
    })

@require_http_methods(['POST'])
def recheck_optimized_resume(request):
    """
    Re-check the edited optimized resume against the job description.
    The frontend sends the edited structured resume as JSON.
    """

    raw = request.POST.get('resume_json') or ''
    jd_text = (request.POST.get('jd') or '').strip()

    # ---------------------------------------------------------
    # Convert JSON from the editor back into Python dictionary
    # ---------------------------------------------------------
    try:
        data = json.loads(raw)
    except (TypeError, ValueError):
        return JsonResponse(
            {
                'error': 'Invalid resume editor data.'
            },
            status=400
        )

    # ---------------------------------------------------------
    # Convert structured resume into plain ATS-readable text
    # ---------------------------------------------------------
    edited_text = optimized_data_to_plain_text(data)

    if len(edited_text.strip()) < min_text_length:
        return JsonResponse(
            {
                'error': 'Please enter more resume content before re-checking.'
            },
            status=400
        )

    # ---------------------------------------------------------
    # Run ATS scoring again
    # ---------------------------------------------------------
    try:
        result = compute_ats_score(
            edited_text,
            jd_text if jd_text else None,
            has_tables=False,
            content_meta={
                'page_count': 1,
                'image_count': 0,
                'significant_images': 0,
                'image_area_ratio': 0.0,
                'is_image_file': False,
            }
        )

    except Exception as exc:
        import traceback
        traceback.print_exc()

        return JsonResponse(
            {
                'error': 'Could not re-check the edited resume.',
                'details': str(exc)
            },
            status=500
        )

    # ---------------------------------------------------------
    # Build actionable JD gaps
    # ---------------------------------------------------------
    jd_gaps = result.get('jd_gaps')

    if jd_gaps is None:
        jd_gaps = build_jd_gaps_from_result(
            result,
            jd_text
        )

    # ---------------------------------------------------------
    # Return new score to JavaScript
    # ---------------------------------------------------------
    return JsonResponse(
        {
            'ats_score': result.get('ats_score', 0),

            'formatting_score':
                result.get('formatting_score', 0),

            'section_score':
                result.get('section_score', 0),

            'keyword_match':
                result.get('keyword_match', 0),

            'missing_sections':
                result.get('missing_sections', []),

            'missing_keywords':
                result.get('missing_keywords', []),

            'jd_gaps':
                jd_gaps,

            'image_based':
                result.get('image_based', False),

            'image_warnings':
                result.get('image_warnings', []),
        }
    )

def optimized_data_to_plain_text(data):
    """Convert the editor structure into ATS-readable text."""
    data = _normalize_ai_resume(data) or {}
    lines = []

    for key in ('name', 'headline'):
        if data.get(key):
            lines.append(data[key])
    if data.get('contact'):
        lines.append(' | '.join(data['contact']))

    if data.get('summary'):
        lines += ['', 'PROFESSIONAL SUMMARY', data['summary']]

    if data.get('skills'):
        lines += ['', 'SKILLS']
        for category, values in data['skills'].items():
            if values:
                lines.append(f"{category}: {', '.join(values)}")

    if data.get('experience'):
        lines += ['', 'EXPERIENCE']
        for item in data['experience']:
            header = ' — '.join(x for x in [item.get('title'), item.get('company')] if x)
            if header:
                lines.append(header)
            meta = ' | '.join(x for x in [item.get('location'), item.get('dates')] if x)
            if meta:
                lines.append(meta)
            for bullet in item.get('bullets', []):
                if bullet:
                    lines.append(f'• {bullet}')

    if data.get('projects'):
        lines += ['', 'PROJECTS']
        for item in data['projects']:
            if item.get('name'):
                lines.append(item['name'])
            if item.get('technologies'):
                lines.append('Technologies: ' + ', '.join(item['technologies']))
            if item.get('description'):
                lines.append(item['description'])
            for bullet in item.get('bullets', []):
                if bullet:
                    lines.append(f'• {bullet}')

    if data.get('education'):
        lines += ['', 'EDUCATION']
        for item in data['education']:
            if item.get('degree'):
                lines.append(item['degree'])
            meta = ' | '.join(x for x in [item.get('institution'), item.get('location'), item.get('dates')] if x)
            if meta:
                lines.append(meta)
            for detail in item.get('details', []):
                if detail:
                    lines.append(f'• {detail}')

    for key, heading in (
        ('certifications', 'CERTIFICATIONS'),
        ('achievements', 'ACHIEVEMENTS'),
        ('languages', 'LANGUAGES'),
    ):
        if data.get(key):
            lines += ['', heading]
            lines.extend(f'• {x}' for x in data[key] if x)

    return '\n'.join(lines).strip()


def _set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        from docx.oxml import OxmlElement
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    from docx.oxml.ns import qn
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def _add_resume_heading(document, text):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = 'Arial'
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '666666')
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def _add_resume_body(document, text, bold=False, italic=False, size=9.5, space_after=2):
    from docx.shared import Pt
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text or ''))
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def _add_resume_bullet(document, text):
    from docx.shared import Pt
    p = document.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Pt(13)
    p.paragraph_format.first_line_indent = Pt(-5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text or ''))
    r.font.name = 'Arial'
    r.font.size = Pt(9.2)
    return p


def build_professional_resume_docx(data):
    """Create a clean, one-column, ATS-readable resume DOCX."""
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    data = _normalize_ai_resume(data) or {}
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    # Header
    if data.get('name'):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(data['name'])
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(18)

    if data.get('headline'):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(data['headline'])
        r.font.name = 'Arial'
        r.font.size = Pt(10.5)
        r.bold = True

    if data.get('contact'):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(' | '.join(data['contact']))
        r.font.name = 'Arial'
        r.font.size = Pt(8.8)

    if data.get('summary'):
        _add_resume_heading(document, 'Professional Summary')
        _add_resume_body(document, data['summary'])

    if data.get('skills'):
        _add_resume_heading(document, 'Skills')
        for category, values in data['skills'].items():
            if not values:
                continue
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1.5)
            r = p.add_run(f'{category}: ')
            r.bold = True
            r.font.name = 'Arial'
            r.font.size = Pt(9.2)
            r2 = p.add_run(', '.join(values))
            r2.font.name = 'Arial'
            r2.font.size = Pt(9.2)

    if data.get('experience'):
        _add_resume_heading(document, 'Experience')
        for item in data['experience']:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            if item.get('title'):
                r = p.add_run(item['title'])
                r.bold = True
            if item.get('company'):
                r = p.add_run((' | ' if item.get('title') else '') + item['company'])
                r.bold = True
            for run in p.runs:
                run.font.name = 'Arial'; run.font.size = Pt(9.5)
            meta = ' | '.join(x for x in [item.get('location'), item.get('dates')] if x)
            if meta:
                _add_resume_body(document, meta, italic=True, size=8.8, space_after=1)
            for bullet in item.get('bullets', []):
                _add_resume_bullet(document, bullet)

    if data.get('projects'):
        _add_resume_heading(document, 'Projects')
        for item in data['projects']:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(item.get('name') or 'Project')
            r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(9.5)
            if item.get('technologies'):
                r2 = p.add_run(' | ' + ', '.join(item['technologies']))
                r2.font.name = 'Arial'; r2.font.size = Pt(8.8)
            if item.get('description'):
                _add_resume_body(document, item['description'], size=9.0, space_after=1)
            for bullet in item.get('bullets', []):
                _add_resume_bullet(document, bullet)

    if data.get('education'):
        _add_resume_heading(document, 'Education')
        for item in data['education']:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(item.get('degree') or 'Qualification')
            r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(9.5)
            meta = ' | '.join(x for x in [item.get('institution'), item.get('location'), item.get('dates')] if x)
            if meta:
                r2 = p.add_run(' | ' + meta)
                r2.font.name = 'Arial'; r2.font.size = Pt(8.8)
            for detail in item.get('details', []):
                _add_resume_bullet(document, detail)

    for key, heading in [('certifications','Certifications'),('achievements','Achievements'),('languages','Languages')]:
        if data.get(key):
            _add_resume_heading(document, heading)
            for value in data[key]:
                _add_resume_bullet(document, value)

    return document

def build_jd_gaps_from_result(result, jd_text):

    if not jd_text:
        return []

    gaps = []

    missing_keywords = (
        result.get('missing_keywords') or []
    )

    for keyword in missing_keywords[:15]:

        keyword = str(keyword).strip()

        if not keyword:
            continue

        lower = keyword.lower()

        # Determine requirement type
        if any(word in lower for word in [
            'experience',
            'candidate',
            'ability',
            'team',
            'company',
            'working',
            'knowledge',
            'strong',
            'years',
            'role'
        ]):
            gap_type = 'Requirement'

        elif any(word in lower for word in [
            'python',
            'java',
            'javascript',
            'typescript',
            'flutter',
            'django',
            'flask',
            'laravel',
            'mysql',
            'postgresql',
            'postgres',
            'sql',
            'mongodb',
            'docker',
            'aws',
            'azure',
            'git',
            'opencv',
            'api',
            'database',
            'framework',
            'machine learning',
            'data'
        ]):
            gap_type = 'Technical skill'

        else:
            gap_type = 'Keyword / concept'

        gaps.append({
            'keyword': keyword,

            'type': gap_type,

            'action': (
                f'Review whether you genuinely have experience '
                f'with "{keyword}". If yes, add it to the most '
                f'relevant Skills, Project, or Experience section '
                f'and describe where you used it. Do not add it '
                f'if you do not.'
            )
        })

    return gaps